"""Generates the comprehensive Review 2 document in DOCX format including Review 1 contents and Review 2 progress."""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for name, w in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{name}')
        node.set(qn('w:w'), str(w))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def create_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles & Colors
    NAVY = RGBColor(27, 54, 93)       # #1B365D - Primary Heading
    DARK_BLUE = RGBColor(41, 70, 115) # #294673 - Subheading
    CHARCOAL = RGBColor(45, 55, 72)   # #2D3748 - Body Text
    MUTED = RGBColor(100, 116, 139)   # #64748B - Captions/Meta

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = CHARCOAL
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # ==========================================
    # TITLE & HEADER
    # ==========================================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("ROBUST DOC")
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    sub_run = sub_p.add_run("Evaluating Structural Robustness of Malicious Document Detectors Under Controlled Mutation")
    sub_run.font.size = Pt(14)
    sub_run.font.bold = True
    sub_run.font.color.rgb = DARK_BLUE

    # Meta Table
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Milestone:", "Review 2 Report (50–60% Implementation & Review 1 Integration)"),
        ("Project Team:", "Jasmeet Ranjan Srivastava (23BCI0019)  |  Radhanjan Neelamraju (23BCI0126)"),
        ("Technical Domain:", "Machine Learning in Cybersecurity / Malware Evasion & Feature Robustness")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        p_k = row.cells[0].paragraphs[0]
        p_v = row.cells[1].paragraphs[0]
        r_k = p_k.add_run(k)
        r_k.bold = True
        r_k.font.color.rgb = NAVY
        r_v = p_v.add_run(v)
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)
        set_cell_background(row.cells[0], "F1F5F9")
        set_cell_background(row.cells[1], "F8FAFC")
        set_cell_margins(row.cells[0], 60, 60, 100, 100)
        set_cell_margins(row.cells[1], 60, 60, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper function for Section Headings
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = DARK_BLUE
        return p

    def add_callout(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.25)
        if bold_prefix:
            r1 = p.add_run(bold_prefix + " ")
            r1.bold = True
            r1.font.color.rgb = NAVY
        r2 = p.add_run(text)
        r2.font.color.rgb = CHARCOAL

    # ==========================================
    # PART I: REVIEW 1 FOUNDATIONS
    # ==========================================
    add_h1("PART I: PROJECT FOUNDATIONS (REVIEW 1 BASELINE)")

    add_h2("1. Problem Statement")
    doc.add_paragraph(
        "PDFs represent one of the most prolific transport mechanisms for weaponized cyber payloads, "
        "leveraging embedded JavaScript, automated launch directives (/OpenAction, /AA), and nested object streams. "
        "Current machine learning-based PDF malware detectors report impressive benchmark accuracies (frequently 98%+), "
        "yet these evaluations are almost exclusively conducted on static, unmodified files."
    )
    doc.add_paragraph(
        "Crucially, standard evaluation metrics fail to reveal two critical operational vulnerabilities:\n"
        "1. How the classifier behaves after a document undergoes routine, benign document-lifecycle actions (e.g., resaving, stream recompression, or metadata editing).\n"
        "2. Which specific features the model's classification confidence actually hinges upon, versus which features break easily under structural shifts."
    )
    add_callout(
        "\"How robust are malicious-document detection methods when the same documents undergo structure-preserving mutations — and which specific features remain reliable versus which are fragile?\"",
        "Core Research Question:"
    )

    add_h2("2. Project Objectives")
    objectives = [
        ("01. Baseline Detector", "Build a standard, representative feature-based PDF classifier using open reference feature sets (PDFRate/Mimicus) with Random Forest classification."),
        ("02. Mutation Engine", "Design a controlled structural mutation engine executing three realistic, non-adversarial perturbations: metadata modification, object reordering, and stream recompression."),
        ("03. Validity & Reachability Gatekeeper", "Construct a strict verification mechanism to prove that every mutated PDF still parses cleanly and that execution triggers retain direct reachability to intact payloads."),
        ("04. Detection Retention Tracking", "Empirically measure the retention percentage of originally detected malware across diverse mutation algorithms."),
        ("05. Per-Feature Stability Analysis", "Track and quantify individual feature volatility to determine which features resist perturbation and which decay rapidly."),
        ("06. Degradation & Tipping-Point Curves", "Parametrize mutation functions with continuous intensity arguments (10% -> 30% -> 60% -> 100%) to locate exact confidence collapse thresholds."),
        ("07. Detector Fragility Report", "Formulate an automated, one-page diagnostic report synthesizing model fragility, feature-level stability, and architectural risks.")
    ]
    for num, desc in objectives:
        p = doc.add_paragraph()
        r1 = p.add_run(f"• {num}: ")
        r1.bold = True
        r1.font.color.rgb = NAVY
        p.add_run(desc)

    add_h2("3. Literature Survey")
    lit_table = doc.add_table(rows=7, cols=3)
    lit_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lit_headers = ["Prior Work", "Methodological Approach", "Identified Research Gap"]
    for c_idx, h in enumerate(lit_headers):
        cell = lit_table.rows[0].cells[c_idx]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, 80, 80, 100, 100)

    lit_rows = [
        ("Smutz & Stavrou (2012) [PDFRate]", "Random Forest classifier evaluated across 202 hand-defined metadata and structural features.", "Reports aggregate accuracy only; feature selection relies on expert intuition without empirical robustness checks."),
        ("Šrndić & Laskov (2014) [Mimicus]", "Constructed targeted evasion attacks against PDFRate with restricted attacker knowledge.", "Investigates mathematically targeted adversarial attacks against a single model rather than routine document noise."),
        ("Xu, Qi & Evans (2016) [EvadeML]", "Applied genetic programming algorithms to mutate PDF structures and evade classifiers while sandbox-verifying payload viability.", "Focuses on mathematically guided evolutionary evasion algorithms rather than realistic everyday document transformations."),
        ("Tong et al. (2019)", "Adversarial retraining using designated 'conserved' features.", "Conserved features were selected through heuristic reasoning rather than experimental feature-level stability profiling."),
        ("Chen et al. (2020)", "Robust classifier training using symbolic interval analysis.", "Defense-centric model hardening that induces elevated false-positive rates on complex commercial documents."),
        ("Liu et al. (CCS 2025)", "PDF object graph combined with language modeling and Graph Neural Networks (GNN).", "Demonstrates broad feature instability under complex representations but does not quantify individual feature-level stability.")
    ]
    for r_idx, (col1, col2, col3) in enumerate(lit_rows):
        row = lit_table.rows[r_idx + 1]
        for c_idx, val in enumerate([col1, col2, col3]):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.add_run(val)
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 90, 90)

    # Set column widths
    for row in lit_table.rows:
        row.cells[0].width = Inches(1.8)
        row.cells[1].width = Inches(2.3)
        row.cells[2].width = Inches(2.4)

    add_h2("4. Research Gap vs. Our Approach")
    gap_table = doc.add_table(rows=5, cols=2)
    gap_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, h in enumerate(["Existing Research Literature", "Our Distinct Diagnostic Approach"]):
        cell = gap_table.rows[0].cells[c_idx]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "294673")
        set_cell_margins(cell, 80, 80, 100, 100)

    gap_data = [
        ("Focus on complex, mathematically optimized adversarial evasion.", "Investigates mild, realistic, non-adversarial structural transformations (routine office processing)."),
        ("Mutations generated via gradient search or genetic algorithms.", "Employs predefined, deterministic structural functions simulating real document workflows."),
        ("Performance reported exclusively as single aggregate accuracy percentages.", "Measures and reports per-feature stability matrices across individual static indicators."),
        ("Binary evaluation criterion (evaded vs. detected).", "Measures continuous probability degradation curves across parameterized mutation intensity levels.")
    ]
    for r_idx, (col1, col2) in enumerate(gap_data):
        row = gap_table.rows[r_idx + 1]
        for c_idx, val in enumerate([col1, col2]):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.add_run(val)
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 90, 90)

    for row in gap_table.rows:
        row.cells[0].width = Inches(3.2)
        row.cells[1].width = Inches(3.3)

    add_h2("5. Architectural Design (Two-Layer Model)")
    doc.add_paragraph(
        "The project is structured into two clean conceptual layers to maintain rigorous experimental separation:\n"
        "• Layer A — Baseline Detector (The Experimental Subject): A deliberately standard, representative feature-based classifier. "
        "No novelty is claimed here; it serves as the test vehicle to ensure findings generalize to commodity detection pipelines.\n"
        "• Layer B — Mutation Engine & Robustness Diagnostic (The Core Contribution): A systematic perturbation suite that executes "
        "controlled structural mutations, enforces payload reachability gatekeeping, and measures feature-by-feature stability."
    )

    add_h2("6. Key Scientific Challenge & The Validity Confound")
    add_callout(
        "If a structural mutation unintentionally corrupts the file or severs the execution trigger (/OpenAction -> /JavaScript), "
        "a subsequent drop in detector confidence merely indicates broken malware, NOT fragile detector features. "
        "To eliminate this experimental confound, our methodology mandates a strict two-stage validity check (V1-V4): "
        "(V1) Syntactic integrity, (V2) Document parser readability, (V3) Trigger structural reachability, and "
        "(V4) Byte-level preservation of the underlying payload stream. Only mutants passing all checks enter evaluation.",
        "The Research Confound & Methodological Control:"
    )

    # ==========================================
    # PART II: REVIEW 2 IMPLEMENTATION PROGRESS
    # ==========================================
    add_h1("PART II: REVIEW 2 IMPLEMENTATION PROGRESS (50–60% MILESTONE)")

    doc.add_paragraph(
        "For the Review 2 milestone, we have transitioned the theoretical architecture into an active, fully functional "
        "experimental testbed. All core modules across Layer A, Layer B, the Safety Gatekeeper, and the Reporting Suite "
        "have been implemented, integrated, and validated across a combined benchmark."
    )

    add_h2("7. Milestone Execution Status Checklist")
    status_table = doc.add_table(rows=7, cols=3)
    status_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, h in enumerate(["Component / Module", "Status", "Implementation Deliverables"]):
        cell = status_table.rows[0].cells[c_idx]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, 80, 80, 100, 100)

    status_data = [
        ("Layer A: Baseline Feature Extractor", "100% Complete", "Extracts 28 structural, keyword, entropy, and metadata features based on PDFRate/Mimicus."),
        ("Layer A: Baseline ML Classifier", "100% Complete", "Random Forest classifier with continuous probabilistic confidence output (scikit-learn)."),
        ("Layer B: Structural Mutation Engine", "100% Complete", "All 3 scoped mutation modules operational with intensity parameters: Metadata, Recompression, Object Reordering."),
        ("Safety & Functionality Gatekeeper", "100% Complete", "Two-stage verifier validating syntax parsing and structural reachability of execution vectors."),
        ("Benchmarking Dataset Corpus", "100% Complete", "Curated 50-PDF testbed: 30 controlled seed files + 20 real-world PDFs (Mozilla & CVE PoCs)."),
        ("Evaluation Suite & Fragility Reporter", "100% Complete", "Automated feature stability tracker, confidence delta calculator, and Fragility Report generator.")
    ]
    for r_idx, (col1, col2, col3) in enumerate(status_data):
        row = status_table.rows[r_idx + 1]
        for c_idx, val in enumerate([col1, col2, col3]):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            if c_idx == 1:
                r.bold = True
                r.font.color.rgb = RGBColor(22, 101, 52)
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)

    for row in status_table.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(1.3)
        row.cells[2].width = Inches(3.0)

    add_h2("8. Detailed Module Implementations")

    p_m1 = doc.add_paragraph()
    r = p_m1.add_run("A. 28-Dimensional Feature Extraction Pipeline (extractor.py)\n")
    r.bold = True
    r.font.color.rgb = NAVY
    p_m1.add_run(
        "Using pikepdf backed by the C++ QPDF parser combined with binary fallback scans, the extractor constructs a "
        "standardized numeric feature vector for each document:\n"
        "• Document Structural Dimensions: file_size, byte_entropy (Shannon calculation, 0.0-8.0), num_pages, num_objects, "
        "num_streams, num_compressed_streams, ratio_compressed_streams, num_indirect_objects, num_trailer_entries.\n"
        "• Execution Triggers & Suspicious Indicators: count_javascript, count_js, count_openaction, count_aa (Additional Actions), "
        "count_acroform, count_launch, count_embedded_files, count_annot, count_uri, count_richmedia.\n"
        "• Object & Graph Architecture: has_xref_stream, has_objstm (Object Stream presence).\n"
        "• Metadata Dictionary Attributes: metadata_length, has_title, has_author, has_creator, has_producer, has_creation_date, has_mod_date."
    )

    p_m2 = doc.add_paragraph()
    r = p_m2.add_run("B. Layer A Baseline Classifier (detector.py)\n")
    r.bold = True
    r.font.color.rgb = NAVY
    p_m2.add_run(
        "Implements a balanced Random Forest model (100 estimators, max_depth=12). Instead of emitting simple binary outputs, "
        "the detector computes continuous calibrated probabilities P(Malicious) in [0.0, 1.0]. This allows our evaluation "
        "to track subtle diagnostic degradation long before a complete prediction flip occurs."
    )

    p_m3 = doc.add_paragraph()
    r = p_m3.add_run("C. Layer B Structural Mutation Engine (mutations/)\n")
    r.bold = True
    r.font.color.rgb = NAVY
    p_m3.add_run(
        "Each mutation is governed by an abstract base interface accepting an intensity parameter lambda in [0.1, 1.0]:\n"
        "1. Metadata Perturbation (metadata_mutation.py): Perturbs or strips document information dictionaries (/Author, /Title, "
        "/Producer, /CreationDate, /ModDate). Simulates routine re-saving in office suites and enterprise sanitization tools.\n"
        "2. Stream Recompression (stream_recompression.py): Adjusts FlateDecode filters across data streams. Selectively compresses, "
        "decompresses, or inverts stream encodings while preserving payload data byte-for-byte. Simulates 'Reduce File Size' optimizations.\n"
        "3. Object Reordering & Linearization (object_reordering.py): Restructures the cross-reference table and re-indexes indirect "
        "objects using Fast Web View linearization algorithms. Alters topological sequence without changing graph semantics."
    )

    p_m4 = doc.add_paragraph()
    r = p_m4.add_run("D. Structural Validity & Reachability Gatekeeper (validator.py)\n")
    r.bold = True
    r.font.color.rgb = NAVY
    p_m4.add_run(
        "Enforces the crucial experimental control. For every candidate mutant, it validates:\n"
        "• V1/V2 Parser Probe: Verifies pikepdf successfully compiles the document root catalog and page tree without fatal syntax errors.\n"
        "• V3/V4 Reachability Check: Validates that if /Root -> /OpenAction, /AA, or /Names -> /JavaScript existed in the original file, "
        "they remain structurally linked and resolve to non-null target action streams in the mutant. Invalid mutants are logged and dropped."
    )

    add_h2("9. Benchmark Dataset Composition (Combined Testbed)")
    doc.add_paragraph(
        "To guarantee high scientific validity, we evaluated the pipeline across both controlled synthetic seeds "
        "and real-world public security samples, totaling 50 PDF documents:"
    )
    data_table = doc.add_table(rows=5, cols=3)
    data_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, h in enumerate(["Dataset Subset", "Count & Distribution", "Structural Characteristics"]):
        cell = data_table.rows[0].cells[c_idx]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, 80, 80, 100, 100)

    data_rows = [
        ("Controlled Synthetic Seeds", "30 PDFs\n(15 Benign, 15 Malicious)", "Synthesized via make_samples.py with deterministic ground truth. Malicious samples feature live /OpenAction JS action streams and /Launch triggers."),
        ("Real-World Benign Corpus", "10 PDFs", "Acquired from the official Mozilla PDF.js test suite. Includes complex multilingual CJK/Arabic encodings, Type3 content streams, font subsets, and Apache PDFBox outputs."),
        ("Real-World Exploit PoCs", "10 PDFs", "Acquired from published security research repositories: CVE-2024-4367 (Arbitrary JS Execution), CVE-2019-14811 (Ghostscript escape), and pentest annotation injections."),
        ("Total Combined Corpus", "50 PDFs\n(25 Benign, 25 Malicious)", "Full benchmark used for training Layer A and conducting the 225-trial structural perturbation battery.")
    ]
    for r_idx, (col1, col2, col3) in enumerate(data_rows):
        row = data_table.rows[r_idx + 1]
        for c_idx, val in enumerate([col1, col2, col3]):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.add_run(val)
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)

    for row in data_table.rows:
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(1.8)
        row.cells[2].width = Inches(2.7)

    # ==========================================
    # PART III: EXPERIMENTAL RESULTS
    # ==========================================
    add_h1("PART III: EXPERIMENTAL RESULTS & EMPIRICAL FINDINGS")

    doc.add_paragraph(
        "We executed the entire end-to-end evaluation suite across all 25 malicious test files. "
        "Each file was subjected to 3 mutation algorithms across 3 intensity levels (lambda = 0.3, 0.6, 1.0), "
        "yielding 225 individual experimental trials and gatekeeper evaluations."
    )

    add_h2("10. Full Experimental Evaluation Matrix (225 Trials)")
    res_table = doc.add_table(rows=4, cols=5)
    res_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, h in enumerate(["Test Cohort", "Total Trials", "Metadata Retention", "Recompression Retention", "Object Reordering Retention"]):
        cell = res_table.rows[0].cells[c_idx]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, 80, 80, 80, 80)

    res_data = [
        ("Real-World CVE Exploit PoCs", "90 trials (10 files)", "100.0%", "100.0%", "96.7% (1 flip)"),
        ("Synthesized Structural Exploits", "135 trials (15 files)", "100.0%", "100.0%", "100.0%"),
        ("Combined Benchmark (Overall)", "225 trials (25 files)", "100.0%", "100.0%", "98.7%")
    ]
    for r_idx, (col1, col2, col3, col4, col5) in enumerate(res_data):
        row = res_table.rows[r_idx + 1]
        for c_idx, val in enumerate([col1, col2, col3, col4, col5]):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            if c_idx >= 2:
                r.bold = True
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 70, 70)

    for row in res_table.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(1.2)
        row.cells[2].width = Inches(1.1)
        row.cells[3].width = Inches(1.1)
        row.cells[4].width = Inches(1.1)

    add_h2("11. Key Diagnostic Discovery: Severe Confidence Degradation")
    doc.add_paragraph(
        "While aggregate binary retention was high (~99%), inspecting the continuous decision boundary "
        "revealed the core discovery of our study: routine structural modifications cause massive diagnostic confidence collapse."
    )

    obs_data = [
        ("CVE-2024-4367 PoC (malicious_real_09):", "Initial confidence: 96.00% -> After Object Reordering (lambda=1.0): Confidence plummeted to 54.00% (-42% confidence drop), bringing the malware to the exact brink of classification failure without touching the payload."),
        ("CVE-2019-14811 PoC (malicious_real_02):", "Initial confidence: 88.00% -> After Object Reordering (lambda=1.0): Confidence dropped to 57.00% (-31% confidence drop). After Stream Recompression (lambda=0.3): Confidence dropped to 58.00% (-30% drop)."),
        ("Synthesized Exploit (malicious_sample_14):", "Initial confidence: 100.00% -> After Metadata Perturbation (lambda=0.6): Confidence dropped to 52.00% (-48% drop). After Stream Recompression: Dropped to 54.00% (-46% drop).")
    ]
    for prefix, body in obs_data:
        p = doc.add_paragraph()
        r = p.add_run(f"• {prefix} ")
        r.bold = True
        r.font.color.rgb = NAVY
        p.add_run(body)

    add_h2("12. Feature Stability Matrix (Empirical Rankings)")
    doc.add_paragraph(
        "Across all 225 trials, we tracked the relative drift of every feature. A feature is deemed stable "
        "if its shift remains within <= 10% of its original baseline value under mutation:"
    )

    feat_table = doc.add_table(rows=11, cols=3)
    feat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, h in enumerate(["Feature Name", "Stability Rate (%)", "Diagnostic Reliability Verdict"]):
        cell = feat_table.rows[0].cells[c_idx]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, 80, 80, 100, 100)

    feat_rows = [
        ("count_javascript / count_js", "100.0%", "Highly Stable — Directly tied to payload triggers."),
        ("count_openaction / count_aa", "100.0%", "Highly Stable — Preserved by semantic necessity."),
        ("count_acroform / count_annot", "100.0%", "Highly Stable — Structural trigger nodes remain intact."),
        ("num_pages", "100.0%", "Highly Stable — Document page topology unaffected."),
        ("byte_entropy", "86.7%", "Moderately Stable — Minor shifts under stream compression."),
        ("num_objects / num_indirect_objects", "54.2%", "Volatile — Shifts during linearization & xref rebuilds."),
        ("num_trailer_entries", "39.6%", "Fragile — Frequently altered during document reserialization."),
        ("file_size", "35.6%", "Fragile — Highly sensitive to stream filter compression."),
        ("num_compressed_streams", "16.4%", "Extremely Fragile — Collapses under standard stream filters."),
        ("ratio_compressed_streams", "15.1%", "Extremely Fragile — Flips drastically during routine resaving.")
    ]
    for r_idx, (col1, col2, col3) in enumerate(feat_rows):
        row = feat_table.rows[r_idx + 1]
        for c_idx, val in enumerate([col1, col2, col3]):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            if c_idx == 1:
                r.bold = True
                pct = float(val.replace("%", ""))
                if pct >= 80:
                    r.font.color.rgb = RGBColor(22, 101, 52)
                elif pct >= 40:
                    r.font.color.rgb = RGBColor(180, 83, 9)
                else:
                    r.font.color.rgb = RGBColor(185, 28, 28)
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)

    for row in feat_table.rows:
        row.cells[0].width = Inches(2.4)
        row.cells[1].width = Inches(1.5)
        row.cells[2].width = Inches(2.6)

    add_h2("13. Output Deliverable: Detector Fragility Report")
    doc.add_paragraph(
        "Our pipeline automatically outputs the one-page diagnostic summary specified in Section 7 of the project brief:"
    )

    report_box = doc.add_paragraph()
    report_box.paragraph_format.left_indent = Inches(0.3)
    report_box.paragraph_format.right_indent = Inches(0.3)
    report_text = (
        "==================================================================\n"
        "                   DETECTOR FRAGILITY REPORT                      \n"
        "==================================================================\n"
        "Model: Random Forest on Combined Benchmark (50 PDFs)\n"
        "Overall Robustness Score: 99.6/100 -> LOW-TO-MODERATE RISK\n"
        "------------------------------------------------------------------\n"
        "Detection Retention per Mutation Type:\n"
        "  * metadata_mutation       : ########## 100.0%\n"
        "  * recompression           : ########## 100.0%\n"
        "  * object_reordering       : #########   98.7%\n"
        "------------------------------------------------------------------\n"
        "Headline Feature Stability (Top Reliable vs Fragile):\n"
        "  [+] High Stability Features:\n"
        "      count_javascript          ########## 100.0%\n"
        "      count_openaction          ########## 100.0%\n"
        "      count_aa                  ########## 100.0%\n"
        "      count_acroform            ########## 100.0%\n"
        "  [-] Fragile / Volatile Features:\n"
        "      num_trailer_entries       ###         39.6%\n"
        "      file_size                 ###         35.6%\n"
        "      num_compressed_streams    #           16.4%\n"
        "      ratio_compressed_streams  #           15.1%\n"
        "------------------------------------------------------------------\n"
        "Verdict: The detector's confidence is heavily influenced by structural & metadata\n"
        "formatting. An attacker or natural document processor can degrade detection\n"
        "confidence by over 40% without modifying core malicious triggers.\n"
        "=================================================================="
    )
    r_rep = report_box.add_run(report_text)
    r_rep.font.name = 'Consolas'
    r_rep.font.size = Pt(8.5)
    r_rep.font.color.rgb = RGBColor(30, 41, 59)

    # ==========================================
    # PART IV: ROADMAP FOR REVIEW 3
    # ==========================================
    add_h1("PART IV: ROADMAP TO FINAL COMPLETION (REVIEW 3)")
    roadmap_items = [
        ("Corpus Expansion (Review 3):", "Scale evaluation corpus to 100+ wild malware samples from the Contagio and VirusShare archives."),
        ("Degradation Curves (Review 3):", "Generate high-resolution continuous line plots plotting detection accuracy against continuous mutation intensity (0% to 100%) to visually depict tipping points."),
        ("SHAP / Explainability Attribution (Stretch Goal):", "Implement SHAP TreeExplainer to attribute exact probability collapse percentages to specific shifting features."),
        ("Defense Blueprint Formulation:", "Synthesize empirical stability metrics into formal feature weighting recommendations for hardening future learning-based classifiers.")
    ]
    for prefix, body in roadmap_items:
        p = doc.add_paragraph()
        r = p.add_run(f"• {prefix} ")
        r.bold = True
        r.font.color.rgb = NAVY
        p.add_run(body)

    # Output path
    output_dir = Path(r"C:\Users\user\Downloads")
    output_path = output_dir / "Robust_Doc_Review_2_Report.docx"
    doc.save(output_path)
    print(f"[+] Successfully generated Review 2 DOCX at: {output_path}")

    # Also copy to workspace
    ws_path = Path(__file__).resolve().parent / "Robust_Doc_Review_2_Report.docx"
    doc.save(ws_path)
    print(f"[+] Saved duplicate in project directory at: {ws_path}")


if __name__ == "__main__":
    create_document()
